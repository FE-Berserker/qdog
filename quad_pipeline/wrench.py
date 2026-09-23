"""quad_pipeline 关节 wrench 提取 (弯矩/剪力/轴向力)。

方法 (verify_sensor.py 已数值验证): 用 MjSpec 在每个关节锚点 (子体局部系原点)
注入 force/torque 传感器, 直接测量父体->子体相互作用 wrench; 传感器输出在子体
局部系, 关节轴恒定, 分解为:
  绕轴分量 = 驱动力矩 (电机/减速器承担)
  垂直分量 = 弯矩     (轴承/壳体承担)
  力       = 轴向 + 剪切
"""

import csv
from pathlib import Path

import mujoco
import numpy as np

from .config import RobotConfig
from .core import run_sim


def build_sensor_model(cfg: RobotConfig, scene_path):
    """加载场景并在 12 关节锚点注入 site + force/torque 传感器。"""
    spec = mujoco.MjSpec.from_file(str(scene_path))
    for lg in cfg.legs:
        for p in cfg.parts:
            jn = cfg.joint_fmt.format(leg=lg, part=p)
            body = spec.joint(jn).parent          # 关节所在子体
            body.add_site(name=f"s_{jn}", pos=[0, 0, 0])
            spec.add_sensor(name=f"frc_{jn}", type=mujoco.mjtSensor.mjSENS_FORCE,
                            objtype=mujoco.mjtObj.mjOBJ_SITE, objname=f"s_{jn}")
            spec.add_sensor(name=f"trq_{jn}", type=mujoco.mjtSensor.mjSENS_TORQUE,
                            objtype=mujoco.mjtObj.mjOBJ_SITE, objname=f"s_{jn}")
    model = spec.compile()
    adr = {}
    for lg in cfg.legs:
        for p in cfg.parts:
            jn = cfg.joint_fmt.format(leg=lg, part=p)
            fi = mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_SENSOR, f"frc_{jn}")
            ti = mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_SENSOR, f"trq_{jn}")
            adr[jn] = (model.sensor_adr[fi], model.sensor_adr[ti])
    return model, adr


def run_wrench(motion_cls, cfg, scene_path):
    """用传感器模型重跑仿真, 返回 wrench 轨迹数组 (time + 12×6)。"""
    model, adr = build_sensor_model(cfg, scene_path)
    sim = motion_cls(cfg, model=model)
    sim.init_pose()
    jn_list = [cfg.joint_fmt.format(leg=lg, part=p) for lg in cfg.legs for p in cfg.parts]
    recs = []
    for _ in range(int(motion_cls.T_END / model.opt.timestep)):
        sim.step()
        row = [sim.data.time]
        for jn in jn_list:
            fa, ta = adr[jn]
            row.extend(float(v) for v in sim.data.sensordata[fa:fa + 3])
            row.extend(float(v) for v in sim.data.sensordata[ta:ta + 3])
        recs.append(row)
        if sim.fallen:
            print(f"!! wrench 重跑跌倒提前终止 t={sim.t:.2f}")
            break
    return np.array(recs)


def decompose(A, cfg):
    """分解为 drive/bend/axial/shear (各 n×12)。"""
    n = A.shape[0]
    out = {k: np.zeros((n, 12)) for k in ("drive", "bend", "axial", "shear")}
    i = 0
    for lg in cfg.legs:
        for p in cfg.parts:
            a = np.array(cfg.part_axis[p], dtype=float)
            f = A[:, 1 + i * 6: 1 + i * 6 + 3]
            nt = A[:, 1 + i * 6 + 3: 1 + i * 6 + 6]
            n_drive = nt @ a
            out["drive"][:, i] = n_drive
            out["bend"][:, i] = np.linalg.norm(nt - np.outer(n_drive, a), axis=1)
            f_ax = f @ a
            out["axial"][:, i] = f_ax
            out["shear"][:, i] = np.linalg.norm(f - np.outer(f_ax, a), axis=1)
            i += 1
    return out


def wrench_csv_head(cfg):
    jn = [cfg.joint_fmt.format(leg=lg, part=p) for lg in cfg.legs for p in cfg.parts]
    return ["time"] + [f"{s}_{n}" for n in jn for s in ("fx", "fy", "fz", "nx", "ny", "nz")]


def save_wrench_csv(path, A, cfg):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(wrench_csv_head(cfg))
        w.writerows([[float(v) for v in row] for row in A])


def make_bending_png(A, dec, cfg, title_cn, out_path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    jn_list = [cfg.joint_fmt.format(leg=lg, part=p) for lg in cfg.legs for p in cfg.parts]
    t = A[:, 0]
    fig, axes = plt.subplots(4, 3, figsize=(15, 11), sharex=True)
    for i in range(12):
        ax = axes[i // 3, i % 3]
        ax.plot(t, dec["bend"][:, i], lw=0.8, color="#a8332a", label="|弯矩|")
        ax.plot(t, np.abs(dec["drive"][:, i]), lw=0.7, ls="--", color="#1f5fa8", label="|绕轴驱动|")
        ax.set_title(jn_list[i], fontsize=9)
        ax.grid(alpha=0.3)
        ax.set_ylabel("N·m", fontsize=8)
        if i == 0:
            ax.legend(fontsize=8, loc="upper right")
    for ax in axes[3]:
        ax.set_xlabel("t (s)")
    fig.suptitle(f"{title_cn}: 关节弯矩 (红) 与绕轴驱动力矩 (蓝虚线)", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    fig.savefig(out_path, dpi=140)
    plt.close(fig)


def write_bending_xlsx(A, dec, cfg, title_cn, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from .report import joint_meta, LEG_CN, PART_CN

    meta = joint_meta(cfg)
    t = A[:, 0]
    wb = Workbook()
    thin = Alignment(horizontal="center", vertical="center")
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)

    ws = wb.active
    ws.title = "关节弯矩时序"
    ws.append([f"{title_cn} — 各关节弯矩 |M_bend| (N·m, 垂直于转轴合力矩), 每 10 ms 抽样"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["时间(s)"] + [m[0].replace("_joint", "") for m in meta])
    for c in ws[3]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, thin
    for i in range(0, len(t), 5):
        ws.append([round(float(t[i]), 3)] + [round(float(v), 2) for v in dec["bend"][i]])
    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = 9
    for j in range(2, 14):
        ws.column_dimensions[get_column_letter(j)].width = 11

    ws2 = wb.create_sheet("关节载荷统计")
    ws2.append([f"{title_cn} — 关节结构载荷统计 (全程)"])
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([])
    ws2.append(["关节", "腿", "部位(转轴)", "|驱动力矩|峰值(N·m)", "弯矩峰值(N·m)", "弯矩RMS(N·m)",
                "弯矩峰值时刻(s)", "|轴向力|峰值(N)", "剪力峰值(N)", "剪力RMS(N)"])
    for c in ws2[3]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, thin
    axis_cn = {"hip": "轴x", "thigh": "轴y", "calf": "轴y"}
    for i, (jn, lg, p, rated) in enumerate(meta):
        ipk = int(np.argmax(dec["bend"][:, i]))
        ws2.append([jn.replace("_joint", ""), LEG_CN[lg], f"{PART_CN[p]}({axis_cn.get(p, '')})",
                    round(float(np.abs(dec["drive"][:, i]).max()), 1),
                    round(float(dec["bend"][:, i].max()), 1),
                    round(float(np.sqrt((dec["bend"][:, i] ** 2).mean())), 2),
                    round(float(t[ipk]), 2),
                    round(float(np.abs(dec["axial"][:, i]).max()), 1),
                    round(float(dec["shear"][:, i].max()), 1),
                    round(float(np.sqrt((dec["shear"][:, i] ** 2).mean())), 1)])
    for j, w in enumerate([16, 7, 16, 17, 14, 13, 14, 13, 11, 11], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "A4"
    wb.properties.creator = "quad_pipeline"
    wb.save(out_path)
