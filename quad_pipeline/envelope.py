"""quad_pipeline 载荷包络: 汇总全部动作的关节 wrench, 输出包络表格与图。"""

import csv
from pathlib import Path

import numpy as np

from .config import RobotConfig
from .wrench import decompose
from .report import joint_meta, LEG_CN, PART_CN

METRICS = ["drive", "bend", "bend_rms", "axial", "shear"]
METRIC_CN = {"drive": "|驱动力矩|峰值(N·m)", "bend": "弯矩峰值(N·m)", "bend_rms": "弯矩RMS(N·m)",
             "axial": "|轴向力|峰值(N)", "shear": "剪力峰值(N)"}


def load_wrench(path, cfg):
    with open(path, encoding="utf-8") as f:
        r = csv.reader(f)
        next(r)
        A = np.array([[float(v) for v in row] for row in r])
    dec = decompose(A, cfg)
    out = {}
    for k in ("drive", "bend", "axial", "shear"):
        out[k] = np.abs(dec[k]).max(axis=0)
    out["bend_rms"] = np.sqrt((dec["bend"] ** 2).mean(axis=0))
    return out


def _safe_save(save_fn, path):
    """输出文件被占用 (如在查看器中打开) 时, 改存 <名>_new.<后缀> 并提示, 不中断流水线。"""
    path = Path(path)
    try:
        save_fn(path)
        return path
    except OSError:
        alt = path.with_name(path.stem + "_new" + path.suffix)
        save_fn(alt)
        print(f"!! {path.name} 被占用 (可能正在查看器中打开), 已改存 {alt.name}; 关闭后可删旧换新")
        return alt


def build_envelope(cfg: RobotConfig, case_wrench_csvs, case_names, case_descs, out_dir):
    """case_wrench_csvs: {key: csv路径}; 输出 包络 xlsx + png。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    meta = joint_meta(cfg)
    keys = [k for k in case_wrench_csvs if Path(case_wrench_csvs[k]).exists()]
    data = {k: load_wrench(case_wrench_csvs[k], cfg) for k in keys}
    nj = len(meta)

    env = {m: np.zeros(nj) for m in METRICS}
    src = {m: [""] * nj for m in METRICS}
    for i in range(nj):
        for m in METRICS:
            vals = [data[k][m][i] for k in keys]
            j = int(np.argmax(vals))
            env[m][i] = vals[j]
            src[m][i] = case_names[keys.index(keys[j])]

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    thin = Alignment(horizontal="center", vertical="center")
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)

    ws = wb.active
    ws.title = "包络汇总"
    ws.append([f"{cfg.name} 关节结构载荷包络 ({len(keys)} 工况取最大)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    hdr = ["关节", "腿", "部位"]
    for m in METRICS:
        hdr += [METRIC_CN[m], "来源工况"]
    ws.append(hdr)
    for c in ws[3]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, thin
    for i, (jn, lg, p, rated) in enumerate(meta):
        row = [jn.replace("_joint", ""), LEG_CN[lg], PART_CN[p]]
        for m in METRICS:
            row += [round(float(env[m][i]), 1), src[m][i]]
        ws.append(row)
    for j, w in enumerate([16, 6, 12] + [13, 9] * len(METRICS), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"

    for m, sheet, unit in (("bend", "弯矩峰值矩阵", "N·m"), ("shear", "剪力峰值矩阵", "N")):
        wsm = wb.create_sheet(sheet)
        wsm.append([f"工况 × 关节 {sheet[:-2]} ({unit})"])
        wsm["A1"].font = Font(bold=True, size=12)
        wsm.append([])
        wsm.append(["工况"] + [mm[0].replace("_joint", "") for mm in meta])
        for c in wsm[3]:
            c.fill, c.font, c.alignment = hdr_fill, hdr_font, thin
        for k in keys:
            wsm.append([case_names[keys.index(k)]] +
                       [round(float(data[k][m][i]), 1) for i in range(nj)])
        wsm.column_dimensions["A"].width = 12
        for j in range(2, 2 + nj):
            wsm.column_dimensions[get_column_letter(j)].width = 10

    ws3 = wb.create_sheet("工况清单")
    ws3.append(["工况", "说明", "wrench 数据文件"])
    for c in ws3[1]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, thin
    for k in keys:
        ws3.append([case_names[keys.index(k)], case_descs[keys.index(k)],
                    str(case_wrench_csvs[k])])
    ws3.column_dimensions["A"].width = 12
    ws3.column_dimensions["B"].width = 44
    ws3.column_dimensions["C"].width = 40

    out_xlsx = out_dir / "关节载荷包络.xlsx"
    wb.properties.creator = "quad_pipeline"
    out_xlsx = _safe_save(wb.save, out_xlsx)

    # ---- 图: 热力图 + 包络柱状 ----
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False

    M = np.array([[data[k]["bend"][i] for i in range(nj)] for k in keys])
    jn_short = [mm[0].replace("_joint", "") for mm in meta]
    fig, axes = plt.subplots(2, 1, figsize=(15, 10), height_ratios=[1, 1.1])
    im = axes[0].imshow(M, aspect="auto", cmap="OrRd")
    axes[0].set_yticks(range(len(keys)), [case_names[keys.index(k)] for k in keys])
    axes[0].set_xticks(range(nj), jn_short, rotation=45, ha="right", fontsize=8)
    for a in range(len(keys)):
        for b in range(nj):
            axes[0].text(b, a, f"{M[a, b]:.0f}", ha="center", va="center", fontsize=7,
                         color="white" if M[a, b] > M.max() * 0.6 else "black")
    axes[0].set_title("关节弯矩峰值热力图 (N·m, 工况×关节)")
    fig.colorbar(im, ax=axes[0], shrink=0.85, label="N·m")

    x = np.arange(nj)
    axes[1].bar(x - 0.2, env["drive"], 0.4, label="|驱动力矩|包络", color="#1f5fa8")
    axes[1].bar(x + 0.2, env["bend"], 0.4, label="弯矩包络", color="#a8332a")
    for i in range(nj):
        axes[1].text(x[i] + 0.2, env["bend"][i] + 2, src["bend"][i], ha="center",
                     fontsize=7, rotation=90, color="#a8332a")
        axes[1].text(x[i] - 0.2, env["drive"][i] + 2, src["drive"][i], ha="center",
                     fontsize=7, rotation=90, color="#1f5fa8")
    axes[1].set_xticks(x, jn_short, rotation=45, ha="right", fontsize=8)
    axes[1].set_ylabel("N·m")
    axes[1].set_title("关节载荷包络 (各工况取最大; 标注=来源工况)")
    axes[1].legend()
    axes[1].grid(alpha=0.3, axis="y")
    axes[1].set_xlim(-0.6, nj - 0.4)
    fig.tight_layout()
    out_png = out_dir / "关节载荷包络.png"
    out_png = _safe_save(lambda p: fig.savefig(p, dpi=140), out_png)
    plt.close(fig)
    return out_xlsx, out_png, env, src
