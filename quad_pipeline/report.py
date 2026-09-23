"""quad_pipeline 报表: 关节扭矩 Excel/曲线 + 运动学曲线。"""

from pathlib import Path

import numpy as np

from .core import C_TQ, C_FZ, C_ST, C_VX, C_VCMD, C_BX, C_BZ, G
from .config import RobotConfig

LEG_CN = {"FL": "左前", "FR": "右前", "RL": "左后", "RR": "右后"}
PART_CN = {"hip": "髋外展(1#)", "thigh": "大腿(髋俯仰)(2#)", "calf": "小腿(膝)(3#)"}
PART_NO = {"hip": "1#", "thigh": "2#", "calf": "3#"}   # 关节序号 (叙述用: 1/2/3号关节)


def joint_meta(cfg: RobotConfig):
    """[(joint_name, leg, part, rated), ...]"""
    out = []
    rated_map = dict(zip(cfg.parts, cfg.rated))
    for lg in cfg.legs:
        for p in cfg.parts:
            out.append((cfg.joint_fmt.format(leg=lg, part=p), lg, p, rated_map[p]))
    return out


def summarize(A, cfg, win):
    t = A[:, 0]
    mask = (t >= win[0]) & (t <= win[1])
    if not mask.any():
        # 仿真提前截断 (跌倒/姿态超限) 落在统计窗口之前: 退化为全部样本
        mask = np.ones(len(t), dtype=bool)
    tq = A[:, C_TQ:C_TQ + 12]
    grf = A[:, C_FZ:C_FZ + 4]
    st = A[:, C_ST:C_ST + 4]
    tq_w, grf_w, st_w = tq[mask], grf[mask], st[mask]
    qw, qx, qy, qz = A[mask, 4], A[mask, 5], A[mask, 6], A[mask, 7]
    roll = np.degrees(np.arctan2(2 * (qw * qx + qy * qz), 1 - 2 * (qx * qx + qy * qy)))
    pitch = np.degrees(np.arcsin(np.clip(2 * (qw * qy - qz * qx), -1.0, 1.0)))
    stance_rms, swing_rms = [], []
    for i in range(12):
        m_st = st_w[:, i // 3] > 0.5
        stance_rms.append(float(np.sqrt(np.mean(tq_w[m_st, i] ** 2))) if m_st.any() else 0.0)
        swing_rms.append(float(np.sqrt(np.mean(tq_w[~m_st, i] ** 2))) if (~m_st).any() else 0.0)
    return {
        "v_mean": float(A[mask, C_VX].mean()),
        "peak_tq": np.abs(tq_w).max(axis=0),
        "rms_tq": np.sqrt((tq_w ** 2).mean(axis=0)),
        "stance_rms": np.array(stance_rms),
        "swing_rms": np.array(swing_rms),
        "grf_peak": float(grf_w.sum(axis=1).max()),
        "grf_mean": float(grf_w.sum(axis=1).mean()),
        "grf_leg_peak": grf_w.max(axis=0),
        "roll_pp": float(roll.max() - roll.min()),
        "pitch_pp": float(pitch.max() - pitch.min()),
        "min_z": float(A[mask, C_BZ].min()),
        "dist": float(A[mask, C_BX][-1] - A[mask, C_BX][0]),
    }


def make_plots(A, cfg, title_cn, out_torques, out_motion, win):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    meta = joint_meta(cfg)
    t = A[:, 0]
    tq = A[:, C_TQ:C_TQ + 12]
    grf = A[:, C_FZ:C_FZ + 4]

    fig, axes = plt.subplots(4, 3, figsize=(15, 11), sharex=True)
    for i, (jn, lg, p, rated) in enumerate(meta):
        ax = axes[i // 3, i % 3]
        ax.plot(t, tq[:, i], lw=0.7, color="#1f5fa8")
        ax.axhline(rated, color="r", ls="--", lw=0.7)
        ax.axhline(-rated, color="r", ls="--", lw=0.7)
        if win:
            ax.axvspan(win[0], min(win[1], t[-1]), color="#2a7a3a", alpha=0.06)
        ax.set_title(f"{jn}({PART_NO[p]})  (额定±{rated:.0f} N·m)", fontsize=9)
        ax.grid(alpha=0.3)
        ax.set_ylabel("N·m", fontsize=8)
    for ax in axes[3]:
        ax.set_xlabel("t (s)")
    fig.suptitle(f"{title_cn}: 各关节执行器力矩", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    fig.savefig(out_torques, dpi=140)
    plt.close(fig)

    qw, qx, qy, qz = A[:, 4], A[:, 5], A[:, 6], A[:, 7]
    roll = np.degrees(np.arctan2(2 * (qw * qx + qy * qz), 1 - 2 * (qx * qx + qy * qy)))
    pitch = np.degrees(np.arcsin(np.clip(2 * (qw * qy - qz * qx), -1.0, 1.0)))
    yaw = np.degrees(np.arctan2(2 * (qw * qz + qx * qy), 1 - 2 * (qy * qy + qz * qz)))

    fig, axes = plt.subplots(2, 3, figsize=(16, 7.5))
    axes[0, 0].plot(t, A[:, C_BX], color="#1f5fa8")
    axes[0, 0].set_title("前进位移 x"); axes[0, 0].set_ylabel("m"); axes[0, 0].grid(alpha=0.3)
    axes[0, 1].plot(t, A[:, C_VX], color="#a8541f", label="实际 vx")
    axes[0, 1].plot(t, A[:, C_VCMD], color="gray", ls="--", label="指令 v_cmd")
    axes[0, 1].set_title("前进速度"); axes[0, 1].set_ylabel("m/s")
    axes[0, 1].legend(fontsize=8); axes[0, 1].grid(alpha=0.3)
    axes[0, 2].plot(t, A[:, C_BZ], color="#1f5fa8")
    axes[0, 2].set_title("躯干高度"); axes[0, 2].set_ylabel("m"); axes[0, 2].grid(alpha=0.3)
    axes[1, 0].plot(t, roll, label="roll", lw=0.8)
    axes[1, 0].plot(t, pitch, label="pitch", lw=0.8)
    axes[1, 0].plot(t, yaw, label="yaw", lw=0.8)
    axes[1, 0].set_title("躯干姿态角"); axes[1, 0].set_ylabel("deg")
    axes[1, 0].set_xlabel("t (s)"); axes[1, 0].legend(fontsize=8); axes[1, 0].grid(alpha=0.3)
    axes[1, 1].plot(t, grf.sum(axis=1), color="#2a7a3a", lw=0.9)
    axes[1, 1].axhline(cfg.mass * G, color="gray", ls="--", lw=0.7, label=f"重力 {cfg.mass * G:.0f} N")
    axes[1, 1].set_title("总垂直地面反力"); axes[1, 1].set_ylabel("N")
    axes[1, 1].set_xlabel("t (s)"); axes[1, 1].legend(fontsize=8); axes[1, 1].grid(alpha=0.3)
    for i, lg in enumerate(cfg.legs):
        axes[1, 2].plot(t, grf[:, i], lw=0.7, label=lg)
    axes[1, 2].set_title("单足垂直反力"); axes[1, 2].set_ylabel("N")
    axes[1, 2].set_xlabel("t (s)"); axes[1, 2].legend(fontsize=8); axes[1, 2].grid(alpha=0.3)
    fig.suptitle(f"{title_cn}: 运动学/地面反力", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(out_motion, dpi=140)
    plt.close(fig)


def write_torque_xlsx(A, cfg, stats, title_cn, desc, win, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    meta = joint_meta(cfg)
    wb = Workbook()
    thin = Alignment(horizontal="center", vertical="center")
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    t = A[:, 0]
    tq = A[:, C_TQ:C_TQ + 12]
    mask = (t >= win[0]) & (t <= win[1])
    if not mask.any():
        mask = np.ones(len(t), dtype=bool)   # 提前截断: 退化为全部样本
    tq_w, t_w = tq[mask], t[mask]

    ws = wb.active
    ws.title = "关节力矩时序"
    ws.append([f"{title_cn} — 各关节执行器力矩 (N·m), 每 10 ms 抽样"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["时间(s)"] + [m[0].replace("_joint", "") for m in meta])
    for c in ws[3]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, thin
    for i in range(0, len(t), 5):
        ws.append([round(float(t[i]), 3)] + [round(float(v), 2) for v in tq[i]])
    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = 9
    for j in range(2, 14):
        ws.column_dimensions[get_column_letter(j)].width = 11

    ws2 = wb.create_sheet("关节力矩统计")
    ws2.append([f"{title_cn} — 关节力矩统计 (窗口 t={win[0]:.1f}~{win[1]:.1f} s)"])
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([])
    ws2.append(["关节", "腿", "部位功能", "额定力矩(N·m)", "正向峰值(N·m)", "负向峰值(N·m)",
                "最大绝对值(N·m)", "RMS(N·m)", "支撑相RMS(N·m)", "摆动相RMS(N·m)",
                "峰值占额定比", "峰值时刻(s)", "备注"])
    for c in ws2[3]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, thin
    for i, (jn, lg, p, rated) in enumerate(meta):
        ipeak = int(np.argmax(np.abs(tq_w[:, i])))
        ws2.append([jn.replace("_joint", ""), LEG_CN[lg], PART_CN[p], rated,
                    round(float(tq_w[:, i].max()), 1), round(float(tq_w[:, i].min()), 1),
                    round(float(stats["peak_tq"][i]), 1), round(float(stats["rms_tq"][i]), 2),
                    round(float(stats["stance_rms"][i]), 2), round(float(stats["swing_rms"][i]), 2),
                    f"{stats['peak_tq'][i] / rated * 100:.0f}%",
                    round(float(t_w[ipeak]), 2),
                    "未超限" if stats["peak_tq"][i] <= rated else "超出额定, 执行器已限幅"])
    for j, w in enumerate([16, 7, 18, 13, 13, 13, 14, 9, 13, 13, 11, 11, 22], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "A4"

    ws3 = wb.create_sheet("工况与关键指标")
    ws3.append(["工况", desc])
    ws3["A1"].font = Font(bold=True, size=12)
    ws3.append([])
    ws3.append(["指标", "数值"])
    for c in ws3[3]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, thin
    rows = [
        ("整机质量 (kg)", round(cfg.mass, 2)),
        ("统计窗口 (s)", f"{win[0]:.2f} ~ {win[1]:.2f}"),
        ("窗口平均前进速度 (m/s)", round(stats["v_mean"], 3)),
        ("窗口前进距离 (m)", round(stats["dist"], 2)),
        ("横滚角峰峰值 (deg)", round(stats["roll_pp"], 2)),
        ("俯仰角峰峰值 (deg)", round(stats["pitch_pp"], 2)),
        ("躯干最低高度 (m)", round(stats["min_z"], 3)),
        ("总地面反力峰值 (N)", round(stats["grf_peak"], 1)),
        ("总地面反力均值 (N)", round(stats["grf_mean"], 1)),
        ("峰值推重比", round(stats["grf_peak"] / (cfg.mass * G), 2)),
        ("单足峰值反力 " + "/".join(cfg.legs) + " (N)",
         " / ".join(f"{v:.0f}" for v in stats["grf_leg_peak"])),
        ("最大关节力矩 (N·m)", round(float(stats["peak_tq"].max()), 1)),
        ("关节力矩峰值部位", meta[int(np.argmax(stats["peak_tq"]))][0]),
    ]
    for k, v in rows:
        ws3.append([k, v])
    ws3.column_dimensions["A"].width = 34
    ws3.column_dimensions["B"].width = 40
    wb.properties.creator = "quad_pipeline"
    wb.save(out_path)
